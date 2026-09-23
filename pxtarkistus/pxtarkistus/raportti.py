"""Raportointi: Excel-työkirja ja tiivis konsolituloste."""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any, Mapping

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

VARIT = {
    "virhe": "FFC7CE",
    "varoitus": "FFEB9C",
    "huomio": "DDEBF7",
    "ok": "E2EFDA",
}
OTSIKKOVARI = "1F4E79"

SARAKENIMET = {
    "tarkistus": "Tarkistus",
    "vakavuus": "Vakavuus",
    "era": "Erätunnus",
    "era_nimi": "Erän nimi",
    "verovuosi": "Verovuosi",
    "tunnusluku": "Tunnusluku",
    "kohde": "Kohde",
    "vertailukohta": "Vertailukohta",
    "odotettu": "Vertailuarvo",
    "saatu": "Havaittu arvo",
    "ero": "Ero",
    "ero_pros": "Ero-%",
    "kuvaus": "Kuvaus",
}


def _muotoile_taulukko(ws, df: pd.DataFrame, vakavuussarake: str | None = "vakavuus") -> None:
    otsikkofontti = Font(bold=True, color="FFFFFF")
    otsikkotausta = PatternFill("solid", fgColor=OTSIKKOVARI)
    reuna = Border(bottom=Side(style="thin", color="B0B0B0"))

    for solu in ws[1]:
        solu.font = otsikkofontti
        solu.fill = otsikkotausta
        solu.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for i, sarake in enumerate(df.columns, start=1):
        kirjain = get_column_letter(i)
        pituus = max([len(str(sarake))] + [len(str(a)) for a in df[sarake].head(200)] + [8])
        ws.column_dimensions[kirjain].width = min(max(pituus + 2, 10), 55)
        if sarake in ("odotettu", "saatu", "ero", "yhteensa", "osasumma", "arvo"):
            for solu in ws[kirjain][1:]:
                solu.number_format = "#,##0"
        if sarake == "ero_pros":
            for solu in ws[kirjain][1:]:
                solu.number_format = "0.0000%"
                if isinstance(solu.value, (int, float)):
                    solu.value = solu.value / 100.0

    if vakavuussarake and vakavuussarake in df.columns:
        sij = list(df.columns).index(vakavuussarake) + 1
        for rivi in range(2, min(ws.max_row, MAX_RIVIT_EXCEL) + 1):
            arvo = ws.cell(row=rivi, column=sij).value
            vari = VARIT.get(str(arvo))
            if vari:
                tausta = PatternFill("solid", fgColor=vari)
                for s in range(1, ws.max_column + 1):
                    ws.cell(row=rivi, column=s).fill = tausta
    for rivi in ws.iter_rows(min_row=2, max_row=min(ws.max_row, MUOTOILU_RIVIT)):
        for solu in rivi:
            solu.border = reuna


def _yhteenveto(havainnot: pd.DataFrame, meta: Mapping[str, Any]) -> pd.DataFrame:
    rivit = [{"Tieto": k, "Arvo": v} for k, v in meta.items()]
    rivit.append({"Tieto": "", "Arvo": ""})
    if havainnot.empty:
        rivit.append({"Tieto": "Havaintoja yhteensä", "Arvo": 0})
        return pd.DataFrame(rivit)
    rivit.append({"Tieto": "Havaintoja yhteensä", "Arvo": int(len(havainnot))})
    for tarkistus, ryhma in havainnot.groupby("tarkistus"):
        for vakavuus, alaryhma in ryhma.groupby("vakavuus"):
            rivit.append({"Tieto": f"{tarkistus} / {vakavuus}", "Arvo": int(len(alaryhma))})
    return pd.DataFrame(rivit)


# Excelissä on noin miljoonan rivin raja, ja satojen tuhansien rivien kirjoitus
# ja muotoilu kestää pitkään. Siksi raportti sisältää vain havainnot ja niiden
# taustaluvut; koko haettu aineisto kirjoitetaan pakattuina CSV-tiedostoina.
MAX_RIVIT_EXCEL = 200_000
MUOTOILU_RIVIT = 5_000


TAULUT_SARAKKEET = {
    "joukko": "Joukko",
    "taulu": "Taulu",
    "otsikko": "Otsikko",
    "paivitetty": "Päivitetty (PxWeb)",
    "data_haettu": "Data haettu",
    "virhe": "Virheitä",
    "varoitus": "Varoituksia",
    "huomio": "Huomioita",
}


def _avain(v: Any) -> str:
    """Vuosi- ym. avain merkkijonoksi: Excelistä luettu 2018.0 -> '2018'."""
    t = str(v)
    return t[:-2] if t.endswith(".0") and t[:-2].lstrip("-").isdigit() else t


def poimi_taustaluvut(havainnot: pd.DataFrame, ristiin: pd.DataFrame | None) -> pd.DataFrame:
    """Näkymävertailun poikkeamista kaikkien taulujen luvut samalle erälle ja vuodelle.

    Näin jokaisen ristiriidan näkee kokonaisuutena – mikä taulu poikkeaa ja
    mitkä ovat keskenään yhtä mieltä – ilman koko aineistoa.
    """
    if ristiin is None or ristiin.empty or havainnot.empty:
        return pd.DataFrame()
    h = havainnot[(havainnot["tarkistus"] == "ristiin") & havainnot["vakavuus"].isin(["virhe", "varoitus"])]
    if h.empty:
        return pd.DataFrame()
    avaimet = {
        (_avain(e), _avain(v), _avain(t))
        for e, v, t in zip(h["era"], h["verovuosi"], h["tunnusluku"])
    }
    r = ristiin
    maski = [
        (_avain(e), _avain(v), _avain(t)) in avaimet
        for e, v, t in zip(r["era"], r["verovuosi"], r["tunnusluku"])
    ]
    poimittu = r[maski]
    sarakkeet = [
        c for c in ("joukko", "era", "era_nimi", "verovuosi", "tunnusluku", "taulu", "arvo", "tila", "paivitetty")
        if c in poimittu
    ]
    return poimittu[sarakkeet].sort_values([c for c in ("era", "verovuosi", "tunnusluku", "taulu") if c in sarakkeet])


def kirjoita_excel(
    tiedosto: Path | str,
    havainnot: pd.DataFrame,
    aineistot: Mapping[str, pd.DataFrame],
    meta: Mapping[str, Any],
    kirjoita_aineisto: bool = True,
) -> Path:
    """Kirjoittaa raportin Exceliksi ja aineiston pakattuina CSV-tiedostoina.

    Palauttaa Excel-tiedoston polun. Aineistot kirjoitetaan samaan kansioon
    nimellä <raportti>_data_<laji>.csv.gz. Jos kirjoita_aineisto on False
    (yhteisraportti, jonka aineisto on jo joukkokohtaisissa tiedostoissa),
    aineistosta käytetään vain ristiriitojen taustaluvut. Aineisto 'taulut'
    (taulujen päivitys- ja hakuajat) kirjoitetaan Taulut-välilehdeksi.
    """
    tiedosto = Path(tiedosto)
    tiedosto.parent.mkdir(parents=True, exist_ok=True)
    taulut = aineistot.get("taulut")

    datatiedostot: list[str] = []
    for nimi, df in aineistot.items():
        if nimi == "taulut":
            continue
        if kirjoita_aineisto and df is not None and not df.empty:
            polku = tiedosto.with_name(f"{tiedosto.stem}_data_{nimi}.csv.gz")
            df.to_csv(polku, index=False, sep=";", encoding="utf-8", compression="gzip")
            datatiedostot.append(polku.name)
    meta = dict(meta)
    if datatiedostot:
        meta["Koko aineisto"] = ", ".join(datatiedostot)

    havainnot_nimetty = havainnot.rename(columns=SARAKENIMET) if not havainnot.empty else havainnot
    if len(havainnot_nimetty) > MAX_RIVIT_EXCEL:
        meta["Huom"] = (
            f"Havaintoja {len(havainnot_nimetty)}; Havainnot-välilehdellä ensimmäiset "
            f"{MAX_RIVIT_EXCEL} vakavuusjärjestyksessä"
        )
    taustaluvut = poimi_taustaluvut(havainnot, aineistot.get("ristiin"))

    with pd.ExcelWriter(tiedosto, engine="openpyxl") as kirjoittaja:
        yhteenveto = _yhteenveto(havainnot, meta)
        yhteenveto.to_excel(kirjoittaja, sheet_name="Yhteenveto", index=False)

        if havainnot.empty:
            pd.DataFrame([{"Tulos": "Ei havaintoja annetuilla toleransseilla"}]).to_excel(
                kirjoittaja, sheet_name="Havainnot", index=False
            )
        else:
            havainnot_nimetty.head(MAX_RIVIT_EXCEL).to_excel(kirjoittaja, sheet_name="Havainnot", index=False)
            for tarkistus, ryhma in havainnot.groupby("tarkistus"):
                nimi = {"ristiin": "Näkymien välillä", "osasummat": "Osasummat",
                        "aikasarja": "Aikasarja", "revisio": "Revisiot",
                        "erahierarkia": "Erähierarkia", "alueet": "Kunnat vs postinumerot"}.get(tarkistus, tarkistus)
                ryhma.rename(columns=SARAKENIMET).head(MAX_RIVIT_EXCEL).to_excel(
                    kirjoittaja, sheet_name=nimi[:31], index=False
                )
        if not taustaluvut.empty:
            taustaluvut.head(MAX_RIVIT_EXCEL).to_excel(
                kirjoittaja, sheet_name="Ristiriitojen taustaluvut", index=False
            )
        if taulut is not None and not taulut.empty:
            taulut.rename(columns=TAULUT_SARAKKEET).to_excel(kirjoittaja, sheet_name="Taulut", index=False)

        kirja = kirjoittaja.book
        for ws in kirja.worksheets:
            if ws.max_row < 1:
                continue
            otsikot = [s.value for s in ws[1]]
            kaannetty = {v: k for k, v in SARAKENIMET.items()}
            alkuperaiset = [kaannetty.get(o, o) for o in otsikot]
            apu = pd.DataFrame(columns=alkuperaiset)
            for i, o in enumerate(alkuperaiset):
                apu[o] = [ws.cell(row=r, column=i + 1).value for r in range(2, min(ws.max_row, 200) + 1)]
            _muotoile_taulukko(ws, apu, "vakavuus" if "vakavuus" in alkuperaiset else None)
    return tiedosto


def _lukumoottori() -> str | None:
    """calamine lukee suuria työkirjoja moninkertaisesti openpyxl:ää nopeammin."""
    try:
        import python_calamine  # noqa: F401
        return "calamine"
    except ImportError:
        return None


def tiivista_raportti(tiedosto: Path | str, ilmoita=None) -> Path:
    """Muuntaa vanhan, aineistovälilehtiä sisältävän raportin kevyeksi.

    Data-välilehdet siirretään pakattuihin CSV-tiedostoihin; muut välilehdet ja
    näkymävertailun ristiriitojen taustaluvut jäävät Exceliin.
    """
    tiedosto = Path(tiedosto)
    try:
        x = pd.ExcelFile(tiedosto, engine=_lukumoottori())
    except (ValueError, ImportError):  # vanha pandas ei tunne calaminea
        x = pd.ExcelFile(tiedosto)
    muut: dict[str, pd.DataFrame] = {}
    aineistot: dict[str, pd.DataFrame] = {}
    for s in x.sheet_names:
        if ilmoita:
            ilmoita(f"  välilehti {s} ...")
        if s.startswith("Data "):
            aineistot[s[5:]] = x.parse(s)
        else:
            muut[s] = x.parse(s)
    kaannetty = {v: k for k, v in SARAKENIMET.items()}
    havainnot = muut.get("Havainnot", pd.DataFrame()).rename(columns=kaannetty)
    if "tarkistus" not in havainnot.columns:
        havainnot = pd.DataFrame()
    elif "ero_pros" in havainnot.columns:
        # Excel tallentaa prosentin murtolukuna (0,05 = 5 %); palautetaan prosenteiksi,
        # ettei uudelleenkirjoitus jaa sitä toiseen kertaan sadalla
        havainnot["ero_pros"] = pd.to_numeric(havainnot["ero_pros"], errors="coerce") * 100
    yv = muut.get("Yhteenveto", pd.DataFrame())
    meta: dict[str, Any] = {}
    for _, rivi in yv.iterrows():
        k = rivi.iloc[0]
        if isinstance(k, str) and k and not k.startswith(("Havaintoja yhteensä",)) and "/" not in k:
            meta[k] = rivi.iloc[1]
    kohde = tiedosto.with_name(f"{tiedosto.stem}_tiivis.xlsx")
    return kirjoita_excel(kohde, havainnot, aineistot, meta)


def konsoliyhteenveto(havainnot: pd.DataFrame, meta: Mapping[str, Any], nayta: int = 12) -> str:
    """Tiivis tekstiyhteenveto ajosta."""
    rivit: list[str] = []
    rivit.append("=" * 78)
    rivit.append("PXTARKISTUS – verotilastojen ristiintarkistus")
    rivit.append("=" * 78)
    for k, v in meta.items():
        rivit.append(f"  {k:<22} {v}")
    rivit.append("-" * 78)
    if havainnot.empty:
        rivit.append("  Ei havaintoja annetuilla toleransseilla.")
        return "\n".join(rivit)

    laskurit = havainnot.groupby(["tarkistus", "vakavuus"]).size()
    for (tarkistus, vakavuus), n in laskurit.items():
        rivit.append(f"  {tarkistus:<12} {vakavuus:<10} {n:>6}")
    rivit.append("-" * 78)

    vakavat = havainnot[havainnot["vakavuus"].isin(["virhe", "varoitus"])].head(nayta)
    if not vakavat.empty:
        rivit.append(f"  Merkittävimmät havainnot (enintään {nayta}):")
        for _, h in vakavat.iterrows():
            nimi = (h.get("era_nimi") or h.get("era") or "")[:38]
            ero = h.get("ero")
            pros = h.get("ero_pros")
            erotieto = ""
            if pd.notna(ero):
                erotieto = f" ero {ero:,.0f}".replace(",", " ")
                if pd.notna(pros):
                    tarkkuus = f"{pros:+.2f} %" if abs(pros) >= 0.01 else f"{pros:+.2e} %"
                    erotieto += f" ({tarkkuus})"
            rivit.append(
                f"   [{h['vakavuus']:<8}] {nimi:<38} {h.get('verovuosi',''):<6}"
                f"{h.get('tunnusluku',''):<6}{erotieto}"
            )
            rivit.append(f"              {h.get('kohde','')}")
    rivit.append("=" * 78)
    return "\n".join(rivit)


def aikaleima() -> str:
    return dt.datetime.now().strftime("%Y-%m-%d_%H%M")
